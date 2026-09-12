# -*- coding: utf-8 -*-
"""Measure every text box with real font metrics and report anything that
spills out of its card, past the footer rule, or off the page.

Boxes are autosize=SHAPE_TO_FIT_TEXT and top-anchored, so the rendered bottom
is top + (wrapped line count * line height).
"""
import sys
from pptx import Presentation
from PIL import ImageFont

DPI = 96.0
FONTS = {
    ("Consolas", False):        r"C:\Windows\Fonts\consola.ttf",
    ("Consolas", True):         r"C:\Windows\Fonts\consolab.ttf",
    ("Segoe UI", False):        r"C:\Windows\Fonts\segoeui.ttf",
    ("Segoe UI", True):         r"C:\Windows\Fonts\segoeuib.ttf",
    ("Segoe UI Light", False):  r"C:\Windows\Fonts\segoeuil.ttf",
    ("Segoe UI Light", True):   r"C:\Windows\Fonts\segoeuisl.ttf",
}
_cache = {}


def font_for(name, size_pt, bold):
    key = (name, bold, round(size_pt, 1))
    if key not in _cache:
        path = FONTS.get((name, bold)) or FONTS[("Segoe UI", False)]
        _cache[key] = ImageFont.truetype(path, int(round(size_pt * DPI / 72.0)))
    return _cache[key]


def wrap_lines(text, font, width_px, spc_px):
    """Greedy word wrap, matching how PowerPoint breaks a paragraph."""
    if not text.strip():
        return 1
    def w(s):
        return font.getlength(s) + spc_px * len(s)
    words, lines, cur = text.split(), 0, ""
    for word in words:
        trial = word if not cur else cur + " " + word
        if w(trial) <= width_px or not cur:
            cur = trial
        else:
            lines += 1
            cur = word
    return lines + (1 if cur else 0)


def measure(shape):
    """Rendered height of a text frame, in inches."""
    tf = shape.text_frame
    width_in = shape.width / 914400.0
    total = 0.0
    for p in tf.paragraphs:
        runs = p.runs
        text = "".join(r.text for r in runs)
        if runs:
            f = runs[0].font
            size = f.size.pt if f.size else 18.0
            name = f.name or "Segoe UI"
            bold = bool(f.bold)
            spc_attr = runs[0]._r.get_or_add_rPr().get('spc')
            spc_px = (int(spc_attr) / 100.0) * DPI / 72.0 if spc_attr else 0.0
        else:
            size, name, bold, spc_px = 18.0, "Segoe UI", False, 0.0
        font = font_for(name, size, bold)
        asc, desc = font.getmetrics()
        line_px = (asc + desc)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        n = wrap_lines(text, font, width_in * DPI, spc_px)
        total += n * line_px * ls / DPI
    return total


def cards(slide):
    """Filled/bordered panels that text is supposed to sit inside."""
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.shape_type is not None \
                and "TEXT_BOX" in str(sh.shape_type):
            continue
        if sh.left is None or sh.width is None:
            continue
        l, t = sh.left / 914400.0, sh.top / 914400.0
        w, h = sh.width / 914400.0, sh.height / 914400.0
        if w >= 2.0 and h >= 1.0 and not (w > 19 and h > 10):   # skip the page ground
            out.append((l, t, w, h))
    return out


def check(path, pad=0.06):
    prs = Presentation(path)
    problems = []
    for i, s in enumerate(prs.slides, 1):
        cs = cards(s)
        for sh in s.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if not sh.text_frame.text.strip():
                continue
            if sh.shape_type is not None and "TEXT_BOX" not in str(sh.shape_type):
                continue
            l, t = sh.left / 914400.0, sh.top / 914400.0
            h = measure(sh)
            bottom = t + h
            label = sh.text_frame.text.replace("\n", " / ")[:46]

            host = None
            for (cl, ct, cw, ch) in cs:
                if cl - 0.02 <= l <= cl + cw and ct - 0.02 <= t <= ct + ch:
                    host = (cl, ct, cw, ch)
                    break
            if host and bottom > host[1] + host[3] - pad:
                problems.append((i, "spills out of its card", label,
                                 round(bottom, 2), round(host[1] + host[3], 2)))
            elif not host and 1.12 < t < 10.28 and bottom > 10.28 - pad:
                problems.append((i, "crosses the footer rule", label,
                                 round(bottom, 2), 10.28))
            elif bottom > 11.25:
                problems.append((i, "runs off the page", label, round(bottom, 2), 11.25))
    return problems


if __name__ == "__main__":
    for path in sys.argv[1:]:
        probs = check(path)
        print("\n==== %s ====" % path)
        if not probs:
            print("  no overflow found")
        for p in probs:
            print("  slide %-3d %-24s bottom %-6s limit %-6s :: %s"
                  % (p[0], p[1], p[3], p[4], p[2]))
