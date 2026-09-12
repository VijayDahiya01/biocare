# -*- coding: utf-8 -*-
"""Content and layout fixes applied to the existing 13 school slides.

Slide numbers below are positions in the ORIGINAL BioCore_for_Schools deck.
"""
from pptx.util import Inches as In, Pt
from pptx.enum.text import PP_ALIGN
from ds import (txt, retext, find, footnote, M, CONTENT_W,
                ACCENT, DEEP, MUTED, INK, SANS, MONO)


def _move(sh, top=None, left=None, height=None, width=None):
    if top is not None:
        sh.top = In(top)
    if left is not None:
        sh.left = In(left)
    if height is not None:
        sh.height = In(height)
    if width is not None:
        sh.width = In(width)


# ---------------------------------------------------------------------------
# Slide 4 - "What you install at the gate"
#   * fixes the body text overflowing the cards and colliding with the footer
#   * demotes the turnstile; the three options are now tablet / CCTV / your gate
# ---------------------------------------------------------------------------
def slide4(s):
    retext(find(s, "A tablet, your existing gate"),
           "A tablet, a camera you own, or the gate you have")
    retext(find(s, "Most schools begin with one tablet"),
           "Most schools begin with one tablet at the main gate. Nothing is drilled, "
           "and nothing at the gate itself has to change.")

    # re-flow: cards start higher and are taller, so nothing spills past them
    for sh in s.shapes:
        if sh.left is None:
            continue
        t = sh.top / 914400.0
        w = sh.width / 914400.0
        is_text = sh.has_text_frame and sh.shape_type is not None \
            and "TEXT_BOX" in str(sh.shape_type)
        if not is_text and abs(w - 6.03) < 0.01 and abs(t - 6.55) < 0.01:
            _move(sh, top=6.20, height=3.40)          # the card itself
        elif not is_text and 6.90 <= t <= 7.60:
            _move(sh, top=t - 0.35)                   # icon artwork
        elif is_text and abs(t - 7.99) < 0.01:
            _move(sh, top=7.57)                       # card title
        elif is_text and abs(t - 8.57) < 0.01:
            _move(sh, top=8.08, height=0.85)          # card body

    bodies = [
        ("A tablet on a stand", [
            "The cheapest way to start. An ordinary tablet on a",
            "floor stand or wall mount, facing the door. No civil",
            "work, no wiring, nothing drilled."],
         "Best for: reception, a side door, a small gate."),
        ("The CCTV you already have", [
            "Most schools already have cameras covering the",
            "gate. Where the angle and the light are good",
            "enough, we read that feed and add nothing at all."],
         "Best for: a gate that is already covered."),
        ("The gate you already own", [
            "Keep the boom barrier, sliding gate or door you",
            "have. Add a camera and a small screen; the software",
            "opens what is already there."],
         "Best for: a working gate that should be smarter."),
    ]
    old_titles = ["A tablet on a stand", "A camera at the gate you have",
                  "A turnstile or flap barrier"]
    for x, (old, (title, body, best)) in zip([1.02, 7.40, 13.79],
                                             zip(old_titles, bodies)):
        tb = find(s, old)
        retext(tb, title)
        body_box = next(sh for sh in s.shapes
                        if sh.has_text_frame and abs(sh.left / 914400.0 - x) < 0.01
                        and abs(sh.top / 914400.0 - 8.08) < 0.01)
        retext(body_box, body)
        txt(s, x, 9.20, 5.19, 0.32, best, 12.5, ACCENT, SANS, ls=1.30)

    retext(find(s, "Existing CCTV cameras can often be reused"),
           "A turnstile or flap barrier is available for controlled college, hostel and "
           "lab areas — at a school gate it only creates the queue you are removing.")
    _move(find(s, "A turnstile or flap barrier is available"), top=9.75)


# ---------------------------------------------------------------------------
# Slide 6 - afternoon. The court-order / watchlist card is replaced by plain
# authorised-guardian verification: same value, none of the custody-liability
# conversation you do not want in a first partnership meeting.
# ---------------------------------------------------------------------------
def slide6(s):
    retext(find(s, "A COURT ORDER"), "15:35  —  A CHANGE OF PLAN")
    retext(find(s, "A barred parent is flagged"), "A one-off is authorised")
    retext(find(s, "Custody disputes are real"), [
        "The office adds an aunt for today only.",
        "The guard sees her the moment she",
        "arrives, and nothing is taken on trust."])


# ---------------------------------------------------------------------------
# Slide 8 - watchlists become something you opt into, not a headline feature.
# ---------------------------------------------------------------------------
def slide8(s):
    retext(find(s, "Alerts only for a named watchlist"),
           "Alerts only where you ask for them")


# ---------------------------------------------------------------------------
# Slide 9 - "Biometric Key Signature" is principal-friendly but not what a CTO
# calls it. Use the real term, and claim only what the implementation does.
# ---------------------------------------------------------------------------
def slide9(s):
    retext(find(s, "The camera measures the face"), [
        "The camera measures the face, turns those measurements into an encrypted",
        "biometric template, and discards the picture. Only the template is kept."])

    label = find(s, "Biometric Key Signature")
    _move(label, width=4.00)
    retext(label, "Encrypted biometric template")
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(13)   # longer term, one line

    caption = find(s, "encrypted — and never a picture")
    _move(caption, width=4.00)
    retext(caption, "encrypted at rest, never an image")


# ---------------------------------------------------------------------------
# Slide 10 - DPDP. Precise rather than sweeping, and consent withdrawal gets
# its own card. The five obligations run along the bottom as a checklist.
# ---------------------------------------------------------------------------
def slide10(s):
    retext(find(s, "India's DPDP Act is stricter"),
           "Designed around the DPDP requirements for children's personal data. The Rules "
           "were notified in November 2025 with phased commencement, so treat this as the "
           "design we build to — not a claim that one slide makes a deployment compliant.")

    retext(find(s, "Enrolment requires a verifiable"), [
        "Enrolment requires a verifiable parental consent,",
        "recorded and timestamped — not a line in a",
        "prospectus nobody read."])

    retext(find(s, "A manual lane always exists"), "Consent can be taken back")
    retext(find(s, "Families who decline are not excluded"), [
        "A parent withdraws and the template is deleted.",
        "Families who decline are never enrolled, and a",
        "non-biometric register runs alongside."])

    retext(find(s, "Arrival and departure only"), [
        "Arrival and departure only — the same events a",
        "paper register holds. No movement history, no",
        "behaviour profile, no stored video."])

    txt(s, M, 9.78, CONTENT_W, 0.36,
        "Verifiable parental consent   ·   Purpose-limited processing   ·   "
        "Consent withdrawal   ·   Non-biometric alternative   ·   No behavioural tracking",
        12.5, DEEP, MONO, ls=1.30, spc=0.6)


# ---------------------------------------------------------------------------
def apply_all(prs, first=0):
    """`first` is the index of original slide 1 inside the assembled deck."""
    sl = prs.slides
    slide4(sl[first + 3])
    slide6(sl[first + 5])
    slide8(sl[first + 7])
    slide9(sl[first + 8])
    slide10(sl[first + 9])


def reframe_school_opener(s):
    """In the partnership deck, slide 1 of the school deck becomes a section start."""
    retext(find(s, "A PLAIN-ENGLISH GUIDE"),
           "THE SCHOOL-FACING DECK  |  WHAT YOUR TEAM SHOWS A PRINCIPAL")
