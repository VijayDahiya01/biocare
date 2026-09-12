# -*- coding: utf-8 -*-
"""Second check: text boxes colliding with each other, or running off the right
margin. verify.py only catches text escaping its card."""
import sys
from pptx import Presentation
from verify import measure, font_for, DPI


def rendered(sh):
    """(left, top, right, bottom) of the ink, not the box."""
    tf = sh.text_frame
    box_w = sh.width / 914400.0
    widest = 0.0
    for p in tf.paragraphs:
        if not p.runs:
            continue
        f = p.runs[0].font
        size = f.size.pt if f.size else 18.0
        font = font_for(f.name or "Segoe UI", size, bool(f.bold))
        spc_attr = p.runs[0]._r.get_or_add_rPr().get('spc')
        spc_px = (int(spc_attr) / 100.0) * DPI / 72.0 if spc_attr else 0.0
        text = "".join(r.text for r in p.runs)
        w = (font.getlength(text) + spc_px * len(text)) / DPI
        widest = max(widest, min(w, box_w))
    l = sh.left / 914400.0
    t = sh.top / 914400.0
    align = tf.paragraphs[0].alignment
    if align is not None and "RIGHT" in str(align):
        return (l + box_w - widest, t, l + box_w, t + measure(sh))
    return (l, t, l + widest, t + measure(sh))


def boxes(slide):
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame or sh.left is None:
            continue
        if sh.shape_type is not None and "TEXT_BOX" not in str(sh.shape_type):
            continue
        if not sh.text_frame.text.strip():
            continue
        out.append((rendered(sh), sh.text_frame.text.replace("\n", " / ")[:40]))
    return out


def check(path, eps=0.035):
    prs = Presentation(path)
    for i, s in enumerate(prs.slides, 1):
        bs = boxes(s)
        for (r, lab) in bs:
            if r[2] > 19.45:
                print("  slide %-3d off right margin  right=%.2f :: %s" % (i, r[2], lab))
            if r[0] < 0.55:
                print("  slide %-3d off left margin   left=%.2f  :: %s" % (i, r[0], lab))
        for a in range(len(bs)):
            for b in range(a + 1, len(bs)):
                (l1, t1, r1, b1), lab1 = bs[a]
                (l2, t2, r2, b2), lab2 = bs[b]
                ox = min(r1, r2) - max(l1, l2)
                oy = min(b1, b2) - max(t1, t2)
                if ox > eps and oy > eps:
                    print("  slide %-3d TEXT OVERLAP %.2fx%.2f :: %r  ~  %r"
                          % (i, ox, oy, lab1, lab2))


if __name__ == "__main__":
    for p in sys.argv[1:]:
        print("\n==== %s ====" % p)
        check(p)
