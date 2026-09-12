# -*- coding: utf-8 -*-
"""BioCore deck design system - extracted from BioCore_for_Schools.pptx."""
import copy
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

# ---- palette (verbatim from the existing deck) ----
BG      = "F5F3FB"   # page ground
INK     = "221C36"   # headline / near-black purple
BODY    = "4A4265"   # deck copy
MUTED   = "6B6480"   # card body
FAINT   = "C4BCDD"   # slide numbers, footer right
RULE    = "DDD8EC"   # hairlines, card borders
ACCENT  = "6D28D9"   # violet
DEEP    = "4C1D95"   # deep violet
TINT    = "EDE7FB"   # accent tint fill
WHITE   = "FFFFFF"
PALE    = "C6BEE2"   # on-dark body
RED     = "C8011A"
REDTINT = "FBE7E9"

# ---- type ----
MONO  = "Consolas"       # eyebrows / labels, uppercase, letter-spaced
LIGHT = "Segoe UI Light" # headlines
SANS  = "Segoe UI"       # everything else

# ---- canvas grid ----
W, H       = 20.0, 11.25
M          = 0.60        # left/right margin
CONTENT_W  = 18.80
COL3       = [0.60, 6.98, 13.37]; COL3_W = 6.03   # 3-up card grid
COL4       = [0.60, 5.39, 10.17, 14.96]; COL4_W = 4.44  # 4-up card grid


def rgb(h):
    return RGBColor.from_string(h)


def txt(slide, l, t, w, h, lines, size, color, font=SANS, bold=False,
        align=PP_ALIGN.LEFT, ls=1.30, spc=None):
    """Text box matching the deck's conventions: no margins, autosize, wrap."""
    box = slide.shapes.add_textbox(In(l), In(t), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = rgb(color)
        if spc:
            r._r.get_or_add_rPr().set('spc', str(int(round(spc * 100))))
    return box


def shape(slide, kind, l, t, w, h, fill=None, line=None, lw=1.0, adj=None, rot=None):
    sh = slide.shapes.add_shape(kind, In(l), In(t), In(w), In(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    if rot is not None:
        sh.rotation = rot
    sh.shadow.inherit = False
    return sh


def card(slide, l, t, w, h, fill=WHITE, line=RULE, lw=1.0, adj=0.05):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill, line, lw, adj)


def bar(slide, l, t, w=2.40, h=0.06, fill=ACCENT):
    return shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill)


def hline(slide, l, t, w, color=RULE, lw=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, In(l), In(t), In(l + w), In(t))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    return c


def vline(slide, l, t, h, color=RULE, lw=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, In(l), In(t), In(l), In(t + h))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    return c


def dot(slide, l, t, d=0.14, fill=ACCENT):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, d, d, fill, adj=0.5)


def arrow(slide, cx, cy, direction="down", size=0.22, fill=ACCENT):
    """Small solid triangle used as a flow marker."""
    rot = {"down": 180, "right": 90, "up": 0, "left": 270}[direction]
    return shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE,
                 cx - size / 2, cy - size / 2, size, size, fill, rot=rot)


def chrome(slide, eyebrow, num, right):
    """Page furniture: ground, eyebrow, number, rules, footer."""
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H, BG)
    txt(slide, M, 0.66, 13.40, 0.34, eyebrow, 11.5, MUTED, MONO, spc=1.8)
    txt(slide, 13.40, 0.66, 6.00, 0.34, num, 11.5, FAINT, MONO,
        align=PP_ALIGN.RIGHT, spc=1.8)
    hline(slide, M, 1.12, CONTENT_W)
    hline(slide, M, 10.28, CONTENT_W)
    txt(slide, M, 10.44, 6.00, 0.32, "BIOCORE", 10.5, MUTED, MONO, spc=1.8)
    txt(slide, 10.40, 10.44, 9.00, 0.32, right, 10.5, FAINT, MONO,
        align=PP_ALIGN.RIGHT, spc=1.4)


def headline(slide, lines, top, size=42, width=15.50, color=INK):
    return txt(slide, M, top, width, 2.60, lines, size, color, LIGHT, ls=1.06)


def deck(slide, lines, top, width=16.20, size=16.5, color=BODY):
    return txt(slide, M, top, width, 1.60, lines, size, color, SANS, ls=1.36)


def footnote(slide, text, top=9.75, color=ACCENT, size=13.5):
    return txt(slide, M, top, CONTENT_W, 0.40, text, size, color, SANS, ls=1.30)


def eyebrow_label(slide, l, t, w, text, color=ACCENT, size=11.0):
    return txt(slide, l, t, w, 0.32, text, size, color, MONO, spc=1.5)


def card_title(slide, l, t, w, text, size=17, color=INK):
    return txt(slide, l, t, w, 0.50, text, size, color, SANS, bold=True, ls=1.14)


def card_body(slide, l, t, w, lines, size=13, color=MUTED):
    return txt(slide, l, t, w, 1.20, lines, size, color, SANS, ls=1.36)


# ---- the vertical spine ----
# A slide lands its headline, rule, deck copy and card band on these lines, so
# the eye does not have to re-find them on every page. Card bodies render at
# about 0.34 in per line at 13 pt, so a three-line body needs
# CARD_T + 2.10 + 1.02 to stay inside its card - which is what sets CARD_H.
H1     = 2.35     # headline top
BAR1   = 3.55     # accent bar, after a one-line headline
DECK1  = 4.00
BAR2   = 4.35     # ...after a two-line headline
DECK2  = 4.80
CARD_T = 6.20     # card band
CARD_H = 3.35
FOOT_Y = 9.80


def icon(slide, x, t, n):
    """Tinted square with a mono numeral - the visual anchor of a card."""
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.42, t + 0.40, 0.82, 0.82,
          fill=TINT, adj=0.18)
    txt(slide, x + 0.42, t + 0.64, 0.82, 0.36, n, 16, ACCENT, MONO,
        align=PP_ALIGN.CENTER, ls=1.0)


def trio(slide, items, t=CARD_T, h=CARD_H):
    """Three cards on the 3-up grid: numeral, title, then body lines."""
    for i, (x, (title, body)) in enumerate(zip(COL3, items), 1):
        card(slide, x, t, COL3_W, h)
        icon(slide, x, t, "%02d" % i)
        card_title(slide, x + 0.42, t + 1.50, 5.19, title)
        card_body(slide, x + 0.42, t + 2.10, 5.19, body)


def quad(slide, items, t=CARD_T, h=CARD_H):
    """Four cards on the 4-up grid."""
    for i, (x, (title, body)) in enumerate(zip(COL4, items), 1):
        card(slide, x, t, COL4_W, h)
        icon(slide, x, t, "%02d" % i)
        card_title(slide, x + 0.42, t + 1.50, 3.60, title, size=16)
        card_body(slide, x + 0.42, t + 2.10, 3.60, body, size=12.5)


def new_deck():
    """A blank presentation on the house canvas."""
    from pptx import Presentation
    prs = Presentation()
    prs.slide_width, prs.slide_height = In(W), In(H)
    return prs


# ---- helpers for editing slides that already exist ----
def find(slide, needle):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise KeyError("no shape containing %r" % needle)


def _set_para_text(p_el, text):
    runs = p_el.findall(qn('a:r'))
    for r in runs[1:]:
        p_el.remove(r)
    runs[0].find(qn('a:t')).text = text


def retext(sh, lines):
    """Replace a text frame's copy, keeping its exact run formatting."""
    if isinstance(lines, str):
        lines = [lines]
    tf = sh.text_frame
    p0 = tf.paragraphs[0]._p
    tmpl = copy.deepcopy(p0)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    _set_para_text(p0, lines[0])
    prev = p0
    for line in lines[1:]:
        np = copy.deepcopy(tmpl)
        _set_para_text(np, line)
        prev.addnext(np)
        prev = np
    return sh


def wipe(slide):
    """Remove every shape, so a slide can be redrawn from scratch."""
    spTree = slide.shapes._spTree
    for sh in list(slide.shapes):
        spTree.remove(sh._element)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def move_to_front(prs, count):
    """Move the last `count` slides to the front, preserving their order."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    tail = ids[-count:]
    for s in tail:
        lst.remove(s)
    for i, s in enumerate(tail):
        lst.insert(i, s)


def arrange(prs, order):
    """Keep only `order` (indices into the current slides), in that sequence."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    keep = [ids[i] for i in order]
    for el in ids:
        if el not in keep:
            prs.part.drop_rel(el.rId)
        lst.remove(el)
    for el in keep:
        lst.append(el)


def renumber(prs, right_footer=None):
    """Rewrite the corner numeral (and optionally the footer) on every slide."""
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            l, t = sh.left / 914400.0, sh.top / 914400.0
            if abs(l - 13.40) < 0.02 and abs(t - 0.66) < 0.02:
                retext(sh, "%02d" % i)
            elif right_footer and abs(l - 10.40) < 0.02 and abs(t - 10.44) < 0.02:
                retext(sh, right_footer)
